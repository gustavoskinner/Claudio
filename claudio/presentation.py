"""
Módulo principal do Claudio - Construtor de apresentações PowerPoint modernas.
"""

from pptx import Presentation as PptxPresentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

from claudio.themes import Theme, THEMES
from claudio.utils import (
    set_slide_background, add_rectangle, add_circle, add_text_box,
    add_accent_bar, add_card, add_decorative_circles, add_bullet_list,
    add_number_highlight, hex_to_rgb, set_shape_fill, set_shape_no_border,
    add_line,
)

# Dimensoes padrao (widescreen 16:9)
SLIDE_WIDTH = Inches(13.33)
SLIDE_HEIGHT = Inches(7.5)


class Presentation:
    """
    Construtor de apresentacoes PowerPoint modernas e sofisticadas.

    Exemplo de uso::

        from claudio import Presentation

        prs = Presentation(theme="tech_startup", titulo="Minha Empresa")
        prs.slide_capa(titulo="Nossa Empresa", subtitulo="Inovacao que transforma")
        prs.slide_agenda(topicos=["Visao Geral", "Produto", "Mercado"])
        prs.slide_conteudo(titulo="Nossa Solucao", pontos=["Escalavel", "Segura"])
        prs.salvar("apresentacao.pptx")
    """

    def __init__(self, theme="corporate_dark", titulo="Apresentacao"):
        if isinstance(theme, str):
            if theme not in THEMES:
                raise ValueError(
                    f"Tema '{theme}' nao encontrado. "
                    f"Disponiveis: {list(THEMES.keys())}"
                )
            self.tema = THEMES[theme]
        else:
            self.tema = theme

        self.prs = PptxPresentation()
        self.prs.slide_width = SLIDE_WIDTH
        self.prs.slide_height = SLIDE_HEIGHT
        self.prs.core_properties.title = titulo
        self.prs.core_properties.author = "Claudio"

        self._blank_layout = self.prs.slide_layouts[6]

    # ------------------------------------------------------------------ slides

    def slide_capa(self, titulo, subtitulo="", autor="", data="",
                   logo_path=None):
        """Slide de capa principal com impacto visual maximo."""
        slide = self._new_slide()
        t = self.tema

        set_slide_background(slide, t.background_rgb)

        # Barra lateral esquerda
        add_rectangle(slide, 0, 0, Inches(0.5), SLIDE_HEIGHT, t.primary_rgb)

        # Blocos decorativos canto superior direito
        add_rectangle(slide, Inches(10.5), 0, Inches(2.83), Inches(2.5),
                      t.primary_rgb, transparency=0.9)
        add_rectangle(slide, Inches(11.5), 0, Inches(1.83), Inches(1.5),
                      t.accent_rgb, transparency=0.7)

        # Circulos decorativos abstratos
        add_decorative_circles(slide, t, count=3)

        # Linha separadora horizontal
        add_rectangle(slide, Inches(0.9), Inches(4.4), Inches(8.0),
                      Inches(0.05), t.primary_rgb)

        # Titulo principal
        self._add_title_text(
            slide, titulo,
            left=Inches(0.9), top=Inches(1.8),
            width=Inches(9.5), height=Inches(2.2),
            font_size=t.size_title, color=t.text_primary_rgb, bold=True
        )

        # Subtitulo
        if subtitulo:
            add_text_box(slide, Inches(0.9), Inches(4.6), Inches(8.5),
                         Inches(0.8), subtitulo, t.font_body, t.size_subtitle,
                         t.accent_rgb)

        # Autor e data
        if autor:
            add_text_box(slide, Inches(0.9), Inches(6.5), Inches(5.0),
                         Inches(0.5), autor, t.font_body, t.size_body,
                         t.text_secondary_rgb)
        if data:
            add_text_box(slide, Inches(8.5), Inches(6.5), Inches(4.0),
                         Inches(0.5), data, t.font_body, t.size_body,
                         t.text_secondary_rgb, alignment=PP_ALIGN.RIGHT)

        if logo_path:
            try:
                slide.shapes.add_picture(logo_path, Inches(10.8), Inches(0.3),
                                         height=Inches(1.0))
            except Exception:
                pass

    def slide_agenda(self, topicos, titulo="Agenda"):
        """Slide de agenda/indice com numeracao moderna."""
        slide = self._new_slide()
        t = self.tema

        set_slide_background(slide, t.background_rgb)
        add_rectangle(slide, 0, 0, SLIDE_WIDTH, Inches(1.6), t.primary_rgb)
        add_rectangle(slide, 0, 0, SLIDE_WIDTH, Inches(0.08), t.accent_rgb)

        self._add_title_text(
            slide, titulo,
            left=Inches(0.6), top=Inches(0.25),
            width=Inches(10.0), height=Inches(1.2),
            font_size=t.size_heading, color=t.text_on_primary_rgb, bold=True
        )

        max_por_col = 5
        colunas = [topicos[:max_por_col], topicos[max_por_col:]] \
            if len(topicos) > max_por_col else [topicos]
        col_starts = [Inches(0.6), Inches(7.0)]

        for col_idx, coluna in enumerate(colunas):
            for i, topico in enumerate(coluna):
                num_global = col_idx * max_por_col + i + 1
                item_top = Inches(2.0) + i * Inches(0.95)
                col_left = col_starts[col_idx]

                add_rectangle(slide, col_left, item_top, Inches(0.55),
                               Inches(0.55), t.primary_rgb, corner_radius=50)

                add_text_box(slide, col_left, item_top, Inches(0.55),
                             Inches(0.6), str(num_global), t.font_title, 16,
                             t.text_on_primary_rgb, bold=True,
                             alignment=PP_ALIGN.CENTER)

                add_text_box(slide, col_left + Inches(0.75), item_top,
                             Inches(5.5), Inches(0.6), topico, t.font_body,
                             t.size_body, t.text_primary_rgb)

        if len(colunas) > 1:
            add_rectangle(slide, Inches(6.3), Inches(1.8), Inches(0.05),
                           Inches(5.5), t.text_secondary_rgb, transparency=0.6)

    def slide_conteudo(self, titulo, texto="", pontos=None, subtitulo=""):
        """Slide de conteudo com titulo, texto e/ou lista de pontos."""
        slide = self._new_slide()
        t = self.tema

        set_slide_background(slide, t.background_rgb)
        self._add_slide_header(slide, titulo, subtitulo)

        content_top = Inches(2.2) if not subtitulo else Inches(2.6)
        content_left = Inches(0.7)
        content_width = Inches(11.9)

        if texto:
            txBox = slide.shapes.add_textbox(
                content_left, content_top, content_width, Inches(1.2))
            txBox.text_frame.word_wrap = True
            p = txBox.text_frame.paragraphs[0]
            p.text = texto
            p.font.name = t.font_body
            p.font.size = Pt(t.size_body)
            p.font.color.rgb = t.text_secondary_rgb
            content_top += Inches(1.3)

        if pontos:
            txBox = slide.shapes.add_textbox(
                content_left, content_top, content_width,
                SLIDE_HEIGHT - content_top - Inches(0.5))
            txBox.text_frame.word_wrap = True
            add_bullet_list(txBox.text_frame, pontos, font_name=t.font_body,
                            font_size=t.size_body, font_color=t.text_primary_rgb,
                            bullet_char="\u25b8")

    def slide_dois_colunas(self, titulo, titulo_esq, conteudo_esq,
                           titulo_dir, conteudo_dir, subtitulo=""):
        """Slide com layout de duas colunas."""
        slide = self._new_slide()
        t = self.tema

        set_slide_background(slide, t.background_rgb)
        self._add_slide_header(slide, titulo, subtitulo)

        col_top = Inches(2.0)
        col_h = Inches(5.0)
        col_w = Inches(5.9)
        gap = Inches(0.35)
        left_x = Inches(0.7)
        right_x = left_x + col_w + gap

        for col_x, col_titulo, conteudo in [
            (left_x, titulo_esq, conteudo_esq),
            (right_x, titulo_dir, conteudo_dir),
        ]:
            add_card(slide, col_x, col_top, col_w, col_h, t, corner_radius=8)
            add_rectangle(slide, col_x, col_top, col_w, Inches(0.55),
                           t.primary_rgb)
            add_text_box(slide, col_x + Inches(0.2), col_top + Inches(0.05),
                         col_w - Inches(0.3), Inches(0.45), col_titulo,
                         t.font_title, 15, t.text_on_primary_rgb, bold=True)

            area_top = col_top + Inches(0.75)
            if isinstance(conteudo, list):
                txBox = slide.shapes.add_textbox(
                    col_x + Inches(0.2), area_top,
                    col_w - Inches(0.4), col_h - Inches(1.0))
                txBox.text_frame.word_wrap = True
                add_bullet_list(txBox.text_frame, conteudo,
                                font_name=t.font_body, font_size=14,
                                font_color=t.text_primary_rgb,
                                bullet_char="\u25b8", spacing=10)
            else:
                add_text_box(slide, col_x + Inches(0.2), area_top,
                             col_w - Inches(0.4), col_h - Inches(1.0),
                             conteudo, t.font_body, 14, t.text_primary_rgb)

    def slide_metricas(self, titulo, metricas, subtitulo=""):
        """Slide de metricas/KPIs em cards com destaque visual."""
        slide = self._new_slide()
        t = self.tema

        set_slide_background(slide, t.background_rgb)
        self._add_slide_header(slide, titulo, subtitulo)

        n = len(metricas)
        if n == 0:
            return

        card_w = Inches(2.8)
        card_h = Inches(2.8)
        spacing = Inches(0.3)
        total_w = n * card_w + (n - 1) * spacing
        start_x = (SLIDE_WIDTH - total_w) / 2
        card_top = Inches(2.7)

        for i, metrica in enumerate(metricas):
            card_left = start_x + i * (card_w + spacing)
            add_card(slide, card_left, card_top, card_w, card_h, t,
                     corner_radius=12)
            cor = t.primary_rgb if i % 2 == 0 else t.accent_rgb
            add_rectangle(slide, card_left, card_top, card_w, Inches(0.12), cor)

            icon_offset = 0
            if metrica.get("icone"):
                add_text_box(slide, card_left, card_top + Inches(0.25),
                             card_w, Inches(0.6), metrica["icone"],
                             t.font_body, 28, cor, alignment=PP_ALIGN.CENTER)
                icon_offset = Inches(0.65)

            add_text_box(slide, card_left, card_top + Inches(0.3) + icon_offset,
                         card_w, Inches(1.0), metrica["valor"], t.font_title,
                         40, cor, bold=True, alignment=PP_ALIGN.CENTER)

            add_text_box(slide, card_left, card_top + Inches(1.5) + icon_offset,
                         card_w, Inches(0.6), metrica["label"], t.font_body,
                         13, t.text_secondary_rgb, alignment=PP_ALIGN.CENTER)

    def slide_citacao(self, citacao, autor="", cargo=""):
        """Slide de citacao/depoimento com design impactante."""
        slide = self._new_slide()
        t = self.tema

        set_slide_background(slide, t.background_rgb)
        add_rectangle(slide, 0, Inches(2.0), SLIDE_WIDTH, Inches(3.5),
                      t.primary_rgb, transparency=0.1)

        add_text_box(slide, Inches(0.5), Inches(1.0), Inches(2.0), Inches(1.5),
                     "\u201c", t.font_title, 96, t.primary_rgb, bold=True)

        self._add_title_text(
            slide, citacao,
            left=Inches(1.2), top=Inches(2.3),
            width=Inches(10.5), height=Inches(2.8),
            font_size=t.size_subtitle, color=t.text_primary_rgb,
            bold=False, alignment=PP_ALIGN.CENTER, italic=True
        )

        if autor:
            add_text_box(slide, Inches(0.5), Inches(5.5), Inches(12.0),
                         Inches(0.5),
                         f"— {autor}" + (f", {cargo}" if cargo else ""),
                         t.font_body, t.size_body, t.accent_rgb, bold=True,
                         alignment=PP_ALIGN.CENTER)

    def slide_timeline(self, titulo, eventos, subtitulo=""):
        """Slide de linha do tempo horizontal."""
        slide = self._new_slide()
        t = self.tema

        set_slide_background(slide, t.background_rgb)
        self._add_slide_header(slide, titulo, subtitulo)

        n = len(eventos)
        if n == 0:
            return

        linha_y = Inches(4.2)
        margem = Inches(0.8)
        area_w = SLIDE_WIDTH - 2 * margem
        passo = area_w / (n - 1) if n > 1 else area_w / 2

        add_rectangle(slide, margem, linha_y - Inches(0.03), area_w,
                      Inches(0.06), t.primary_rgb)

        for i, evento in enumerate(eventos):
            x = margem + i * passo if n > 1 else margem + area_w / 2
            dot_size = Inches(0.35)
            cor = t.primary_rgb if i % 2 == 0 else t.accent_rgb
            add_circle(slide, x - dot_size / 2, linha_y - dot_size / 2,
                       dot_size, cor)

            if i % 2 == 0:
                ano_y = linha_y - Inches(1.2)
            else:
                ano_y = linha_y + Inches(0.5)

            add_text_box(slide, x - Inches(0.8), ano_y, Inches(1.6),
                         Inches(0.45), evento.get("ano", ""), t.font_title,
                         16, cor, bold=True, alignment=PP_ALIGN.CENTER)
            add_text_box(slide, x - Inches(0.9), ano_y + Inches(0.45),
                         Inches(1.8), Inches(0.45), evento.get("titulo", ""),
                         t.font_body, 13, t.text_primary_rgb, bold=True,
                         alignment=PP_ALIGN.CENTER)
            if evento.get("descricao"):
                add_text_box(slide, x - Inches(0.9), ano_y + Inches(0.9),
                             Inches(1.8), Inches(0.8), evento["descricao"],
                             t.font_body, 11, t.text_secondary_rgb,
                             alignment=PP_ALIGN.CENTER)

    def slide_divisor(self, titulo_secao, numero="", descricao=""):
        """Slide divisor de secao com design impactante."""
        slide = self._new_slide()
        t = self.tema

        set_slide_background(slide, t.primary_rgb)
        add_circle(slide, Inches(7.5), Inches(-1.5), Inches(6.0),
                   t.secondary_rgb, transparency=0.6)
        add_circle(slide, Inches(-1.5), Inches(4.0), Inches(4.0),
                   t.secondary_rgb, transparency=0.7)
        add_circle(slide, Inches(9.0), Inches(4.5), Inches(3.0),
                   t.accent_rgb, transparency=0.8)

        if numero:
            add_text_box(slide, Inches(0.8), Inches(1.5), Inches(3.0),
                         Inches(2.5), numero, t.font_title, 96,
                         t.text_on_primary_rgb, bold=True)
            self._add_title_text(
                slide, titulo_secao,
                left=Inches(3.5), top=Inches(2.8),
                width=Inches(8.0), height=Inches(0.7),
                font_size=t.size_heading, color=t.text_on_primary_rgb,
                bold=True
            )
        else:
            self._add_title_text(
                slide, titulo_secao,
                left=Inches(0.8), top=Inches(2.5),
                width=Inches(11.0), height=Inches(2.0),
                font_size=t.size_title, color=t.text_on_primary_rgb,
                bold=True
            )

        if descricao:
            add_text_box(
                slide,
                Inches(3.5) if numero else Inches(0.8), Inches(4.5),
                Inches(8.5), Inches(0.8),
                descricao, t.font_body, t.size_body, t.text_on_primary_rgb
            )

        add_rectangle(slide, 0, Inches(6.8), Inches(4.0), Inches(0.12),
                      t.accent_rgb)

    def slide_comparacao(self, titulo, opcao_a, opcao_b, subtitulo=""):
        """Slide de comparacao entre duas opcoes."""
        slide = self._new_slide()
        t = self.tema

        set_slide_background(slide, t.background_rgb)
        self._add_slide_header(slide, titulo, subtitulo)

        card_top = Inches(2.0)
        card_h = Inches(5.0)
        card_w = Inches(5.9)

        for idx, (opcao, x) in enumerate([
            (opcao_a, Inches(0.4)),
            (opcao_b, Inches(6.9))
        ]):
            cor_hex = opcao.get("cor", t.primary if idx == 0 else t.accent)
            cor = t.rgb(cor_hex)
            add_card(slide, x, card_top, card_w, card_h, t, corner_radius=10)
            add_rectangle(slide, x, card_top, card_w, Inches(0.65), cor)
            add_text_box(slide, x + Inches(0.2), card_top + Inches(0.1),
                         card_w - Inches(0.3), Inches(0.5),
                         opcao.get("titulo", ""), t.font_title, 18,
                         t.text_on_primary_rgb, bold=True)

            pontos = opcao.get("pontos", [])
            if pontos:
                txBox = slide.shapes.add_textbox(
                    x + Inches(0.25), card_top + Inches(0.85),
                    card_w - Inches(0.5), card_h - Inches(1.1))
                txBox.text_frame.word_wrap = True
                add_bullet_list(txBox.text_frame, pontos, font_name=t.font_body,
                                font_size=14, font_color=t.text_primary_rgb,
                                bullet_char="\u2713" if idx == 1 else "\u25b8",
                                spacing=10)

        add_text_box(slide, Inches(6.1), Inches(4.0), Inches(1.1), Inches(0.6),
                     "VS", t.font_title, 22, t.text_secondary_rgb, bold=True,
                     alignment=PP_ALIGN.CENTER)

    def slide_encerramento(self, mensagem="Obrigado!", contato="",
                           website="", subtexto=""):
        """Slide de encerramento com design impactante."""
        slide = self._new_slide()
        t = self.tema

        set_slide_background(slide, t.background_rgb)
        add_rectangle(slide, 0, Inches(3.8), SLIDE_WIDTH, Inches(3.7),
                      t.primary_rgb)
        add_circle(slide, Inches(9.5), Inches(-0.5), Inches(3.0),
                   t.accent_rgb, transparency=0.8)
        add_circle(slide, Inches(-1.0), Inches(5.0), Inches(2.5),
                   t.secondary_rgb, transparency=0.7)
        add_rectangle(slide, Inches(4.0), Inches(3.65), Inches(5.33),
                      Inches(0.08), t.accent_rgb)

        self._add_title_text(
            slide, mensagem,
            left=Inches(0.5), top=Inches(1.5),
            width=Inches(12.33), height=Inches(2.0),
            font_size=t.size_title + 4, color=t.text_primary_rgb, bold=True,
            alignment=PP_ALIGN.CENTER
        )

        if subtexto:
            add_text_box(slide, Inches(0.5), Inches(3.0), Inches(12.33),
                         Inches(0.6), subtexto, t.font_body, t.size_body,
                         t.text_secondary_rgb, alignment=PP_ALIGN.CENTER)

        info_y = Inches(4.3)
        if contato:
            add_text_box(slide, Inches(0.5), info_y, Inches(12.33), Inches(0.55),
                         contato, t.font_body, t.size_subtitle,
                         t.text_on_primary_rgb, bold=True,
                         alignment=PP_ALIGN.CENTER)
            info_y += Inches(0.7)

        if website:
            add_text_box(slide, Inches(0.5), info_y, Inches(12.33), Inches(0.5),
                         website, t.font_body, t.size_body,
                         t.text_on_primary_rgb, alignment=PP_ALIGN.CENTER)

    # ------------------------------------------------------------- internos

    def _new_slide(self):
        return self.prs.slides.add_slide(self._blank_layout)

    def _add_slide_header(self, slide, titulo, subtitulo=""):
        t = self.tema
        header_h = Inches(1.5) if subtitulo else Inches(1.3)
        add_rectangle(slide, 0, 0, SLIDE_WIDTH, header_h, t.primary_rgb)
        add_rectangle(slide, 0, 0, SLIDE_WIDTH, Inches(0.08), t.accent_rgb)

        add_text_box(slide, Inches(0.6), Inches(0.18), Inches(12.0),
                     Inches(0.9), titulo, t.font_title, t.size_heading,
                     t.text_on_primary_rgb, bold=True)

        if subtitulo:
            add_text_box(slide, Inches(0.6), Inches(0.98), Inches(12.0),
                         Inches(0.45), subtitulo, t.font_body, t.size_caption,
                         t.text_on_primary_rgb)

        num = len(self.prs.slides)
        add_text_box(slide, Inches(12.3), Inches(7.15), Inches(0.9),
                     Inches(0.3), str(num), t.font_body, t.size_small,
                     t.text_secondary_rgb, alignment=PP_ALIGN.RIGHT)

    def _add_title_text(self, slide, texto, left, top, width, height,
                        font_size, color, bold=True,
                        alignment=PP_ALIGN.LEFT, italic=False):
        t = self.tema
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.word_wrap = True
        tf.auto_size = None

        p = tf.paragraphs[0]
        p.text = texto
        p.font.name = t.font_title
        p.font.size = Pt(font_size)
        p.font.bold = bold
        p.font.italic = italic
        p.font.color.rgb = color
        p.alignment = alignment
        p.space_after = Pt(0)
        p.space_before = Pt(0)
        return txBox

    def salvar(self, caminho):
        """Salva a apresentacao no caminho especificado (.pptx)."""
        import os
        if not caminho.endswith(".pptx"):
            caminho += ".pptx"
        self.prs.save(caminho)
        return os.path.abspath(caminho)

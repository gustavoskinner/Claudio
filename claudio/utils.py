"""
Utilitários para criação de elementos visuais sofisticados.

Formas geométricas, barras decorativas, cards e outros elementos modernos.
"""

from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

from claudio.themes import Theme


def hex_to_rgb(hex_color: str) -> RGBColor:
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def set_shape_fill(shape, color: RGBColor, transparency: float = 0.0):
    """Define o preenchimento sólido de uma forma."""
    fill = shape.fill
    fill.solid()
    fill.fore_color.rgb = color
    if transparency > 0:
        sp = shape._element
        sp_pr = sp.find(qn("p:spPr"))
        if sp_pr is not None:
            a_solid = sp_pr.find(qn("a:solidFill"))
            if a_solid is not None:
                srgb = a_solid.find(qn("a:srgbClr"))
                if srgb is not None:
                    alpha = srgb.makeelement(qn("a:alpha"), {})
                    alpha.set("val", str(int((1 - transparency) * 100000)))
                    srgb.append(alpha)


def set_shape_no_border(shape):
    """Remove a borda de uma forma."""
    shape.line.fill.background()


def add_rectangle(slide, left, top, width, height, color: RGBColor,
                  transparency: float = 0.0, corner_radius: int = 0):
    """Adiciona um retângulo com opções de estilo."""
    if corner_radius > 0:
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
        )
        # Ajustar raio do canto
        shape.adjustments[0] = corner_radius / 100.0
    else:
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, left, top, width, height
        )
    set_shape_fill(shape, color, transparency)
    set_shape_no_border(shape)
    return shape


def add_circle(slide, left, top, size, color: RGBColor, transparency: float = 0.0):
    """Adiciona um círculo decorativo."""
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    set_shape_fill(shape, color, transparency)
    set_shape_no_border(shape)
    return shape


def add_line(slide, start_x, start_y, end_x, end_y, color: RGBColor,
             width: float = 1.0):
    """Adiciona uma linha decorativa."""
    connector = slide.shapes.add_connector(
        1, start_x, start_y, end_x, end_y  # MSO_CONNECTOR.STRAIGHT
    )
    connector.line.color.rgb = color
    connector.line.width = Pt(width)
    return connector


def add_accent_bar(slide, left, top, width, height, theme: Theme):
    """Adiciona uma barra de destaque com a cor primária do tema."""
    return add_rectangle(slide, left, top, width, height, theme.primary_rgb)


def add_card(slide, left, top, width, height, theme: Theme,
             corner_radius: int = 8, shadow: bool = True):
    """Adiciona um card com estilo moderno."""
    card = add_rectangle(
        slide, left, top, width, height,
        theme.surface_rgb, corner_radius=corner_radius
    )
    if shadow:
        _add_shadow(card)
    return card


def _add_shadow(shape):
    """Adiciona sombra suave a uma forma via XML."""
    sp = shape._element
    sp_pr = sp.find(qn("a:effectLst"))
    if sp_pr is None:
        effect_lst = sp.makeelement(qn("a:effectLst"), {})
        # Sombra externa suave
        outer_shdw = effect_lst.makeelement(qn("a:outerShdw"), {
            "blurRad": "76200",
            "dist": "38100",
            "dir": "5400000",
            "algn": "bl",
            "rotWithShape": "0",
        })
        srgb_clr = outer_shdw.makeelement(qn("a:srgbClr"), {"val": "000000"})
        alpha = srgb_clr.makeelement(qn("a:alpha"), {"val": "23000"})
        srgb_clr.append(alpha)
        outer_shdw.append(srgb_clr)
        effect_lst.append(outer_shdw)
        sp_pr_element = sp.find(qn("p:spPr"))
        if sp_pr_element is not None:
            sp_pr_element.append(effect_lst)


def add_text_box(slide, left, top, width, height, text: str,
                 font_name: str = "Calibri", font_size: int = 16,
                 font_color: RGBColor = None, bold: bool = False,
                 alignment: PP_ALIGN = PP_ALIGN.LEFT,
                 vertical_anchor: MSO_ANCHOR = MSO_ANCHOR.TOP):
    """Adiciona uma caixa de texto estilizada."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    txBox.text_frame.word_wrap = True
    txBox.text_frame.auto_size = None

    p = txBox.text_frame.paragraphs[0]
    p.text = text
    p.alignment = alignment
    p.font.name = font_name
    p.font.size = Pt(font_size)
    p.font.bold = bold
    if font_color:
        p.font.color.rgb = font_color

    txBox.text_frame.paragraphs[0].space_after = Pt(0)
    txBox.text_frame.paragraphs[0].space_before = Pt(0)

    return txBox


def add_icon_placeholder(slide, left, top, size, icon_text: str,
                         theme: Theme, style: str = "filled"):
    """Adiciona um placeholder de ícone (círculo com texto/emoji)."""
    if style == "filled":
        circle = add_circle(slide, left, top, size, theme.primary_rgb)
        text_color = theme.text_on_primary_rgb
    elif style == "outline":
        circle = add_circle(slide, left, top, size, theme.background_rgb, transparency=0.9)
        circle.line.color.rgb = theme.primary_rgb
        circle.line.width = Pt(2)
        text_color = theme.primary_rgb
    else:
        circle = add_circle(slide, left, top, size, theme.accent_rgb)
        text_color = theme.text_on_primary_rgb

    tf = circle.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = icon_text
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(int(size / Inches(1) * 18))
    p.font.color.rgb = text_color
    tf.paragraphs[0].space_after = Pt(0)
    tf.paragraphs[0].space_before = Pt(0)

    circle.text_frame.auto_size = None
    try:
        circle.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        circle.text_frame.word_wrap = False
    except Exception:
        pass

    return circle


def add_decorative_circles(slide, theme: Theme, count: int = 3):
    """Adiciona círculos decorativos abstratos ao fundo do slide."""
    import random
    positions = [
        (Inches(8.5), Inches(-0.5), Inches(2.5)),
        (Inches(-0.8), Inches(5.5), Inches(2.0)),
        (Inches(9.0), Inches(4.5), Inches(1.8)),
        (Inches(7.5), Inches(-1.0), Inches(3.0)),
        (Inches(-1.0), Inches(-0.5), Inches(2.2)),
    ]
    colors = [theme.primary_rgb, theme.secondary_rgb, theme.accent_rgb]

    for i in range(min(count, len(positions))):
        left, top, size = positions[i]
        color = colors[i % len(colors)]
        add_circle(slide, left, top, size, color, transparency=0.85)


def add_progress_bar(slide, left, top, width, height, progress: float,
                     theme: Theme):
    """Adiciona uma barra de progresso."""
    # Fundo
    add_rectangle(slide, left, top, width, height,
                  theme.text_secondary_rgb, transparency=0.7)
    # Preenchimento
    fill_width = int(width * min(max(progress, 0), 1.0))
    if fill_width > 0:
        add_rectangle(slide, left, top, fill_width, height,
                      theme.primary_rgb, corner_radius=0)


def add_number_highlight(slide, left, top, number: str, label: str,
                         theme: Theme, size: str = "large"):
    """Adiciona um número em destaque com label (para KPIs/métricas)."""
    if size == "large":
        num_size = 48
        label_size = 14
        box_w = Inches(2.5)
        box_h = Inches(1.8)
    else:
        num_size = 36
        label_size = 12
        box_w = Inches(2.0)
        box_h = Inches(1.4)

    # Card de fundo
    card = add_card(slide, left, top, box_w, box_h, theme, corner_radius=10)

    # Número
    add_text_box(
        slide, left + Inches(0.15), top + Inches(0.2),
        box_w - Inches(0.3), Inches(0.8),
        number, theme.font_title, num_size,
        theme.primary_rgb, bold=True, alignment=PP_ALIGN.CENTER
    )

    # Label
    add_text_box(
        slide, left + Inches(0.15), top + Inches(0.95),
        box_w - Inches(0.3), Inches(0.5),
        label, theme.font_body, label_size,
        theme.text_secondary_rgb, alignment=PP_ALIGN.CENTER
    )

    return card


def add_bullet_list(text_frame, items: list, font_name: str = "Calibri",
                    font_size: int = 16, font_color: RGBColor = None,
                    bullet_char: str = "\u2022", spacing: int = 12):
    """Adiciona uma lista com marcadores estilizados a um text frame."""
    for i, item in enumerate(items):
        if i == 0:
            p = text_frame.paragraphs[0]
        else:
            p = text_frame.add_paragraph()

        p.text = f"{bullet_char}  {item}"
        p.font.name = font_name
        p.font.size = Pt(font_size)
        if font_color:
            p.font.color.rgb = font_color
        p.space_after = Pt(spacing)
        p.space_before = Pt(0)
        p.level = 0


def set_slide_background(slide, color: RGBColor):
    """Define a cor de fundo de um slide."""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color
